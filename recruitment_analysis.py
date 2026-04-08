# recruitment_analysis.py
# Run with:
#   python -m streamlit run recruitment_analysis.py

from __future__ import annotations

from io import BytesIO
from datetime import datetime
import re
from typing import Dict, List, Tuple

import pandas as pd
import plotly.express as px
import streamlit as st

# ── Frontend systems ──────────────────────────────────────────────────────
from themes import ThemeFactory
from canvas_components import (
    render_particle_network,
    render_wave_divider,
    render_kpi_ring,
    render_mosaic_art,
)
from web_artifacts import generate_html_report, generate_chart_artifact

# Optional PPT export
try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except Exception:
    PPTX_AVAILABLE = False


REQUIRED_COLUMNS = [
    "Recruiter",
    "Joining_Date",
    "Location",
    "Client",
    "Candidate",
    "Manager",
    "Work_Authorization",
]


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value).strip())


def detect_name_variants(series: pd.Series) -> List[List[str]]:
    names = (
        series.dropna()
        .astype(str)
        .map(normalize_text)
        .replace("", pd.NA)
        .dropna()
        .unique()
        .tolist()
    )
    buckets: Dict[str, List[str]] = {}
    for name in names:
        key = re.sub(r"[^a-z]", "", name.lower())
        key = key[:10] if key else "unknown"
        buckets.setdefault(key, []).append(name)
    return [sorted(v) for v in buckets.values() if len(v) > 1]


def preprocess(df: pd.DataFrame) -> pd.DataFrame:
    data = df.copy()

    missing = [c for c in REQUIRED_COLUMNS if c not in data.columns]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    for col in ["Recruiter", "Location", "Client", "Candidate", "Manager", "Work_Authorization"]:
        data[col] = data[col].astype(str).map(normalize_text)

    data["Joining_Date"] = pd.to_datetime(data["Joining_Date"], errors="coerce")
    data = data.dropna(subset=["Joining_Date"]).copy()

    data["Month"] = data["Joining_Date"].dt.to_period("M").dt.to_timestamp()
    data["Weekday"] = data["Joining_Date"].dt.day_name()
    data["Work_Authorization"] = data["Work_Authorization"].str.upper()
    data["Work_Mode"] = data["Location"].str.contains(
        r"hybrid|remote", case=False, na=False
    ).map({True: "Hybrid/Remote", False: "Onsite/Other"})

    return data


def apply_filters(df: pd.DataFrame) -> pd.DataFrame:
    st.sidebar.header("Filters")

    min_date = df["Joining_Date"].min().date()
    max_date = df["Joining_Date"].max().date()

    selected_range = st.sidebar.date_input(
        "Joining date range",
        value=(min_date, max_date),
        min_value=min_date,
        max_value=max_date,
    )
    if isinstance(selected_range, tuple) and len(selected_range) == 2:
        start_date, end_date = selected_range
    else:
        start_date, end_date = min_date, max_date

    recruiters = sorted(df["Recruiter"].dropna().unique().tolist())
    clients = sorted(df["Client"].dropna().unique().tolist())
    managers = sorted(df["Manager"].dropna().unique().tolist())

    selected_recruiters = st.sidebar.multiselect("Recruiter", recruiters, default=recruiters)
    selected_clients = st.sidebar.multiselect("Client", clients, default=clients)
    selected_managers = st.sidebar.multiselect("Manager", managers, default=managers)

    filtered = df[
        (df["Joining_Date"].dt.date >= start_date)
        & (df["Joining_Date"].dt.date <= end_date)
        & (df["Recruiter"].isin(selected_recruiters))
        & (df["Client"].isin(selected_clients))
        & (df["Manager"].isin(selected_managers))
    ].copy()

    return filtered


def generate_insights(df: pd.DataFrame) -> List[str]:
    if df.empty:
        return ["No records match current filters."]

    total = len(df)
    unique_clients = df["Client"].nunique()
    unique_recruiters = df["Recruiter"].nunique()

    top_recruiter_counts = df["Recruiter"].value_counts()
    top_client_counts = df["Client"].value_counts()
    top_manager_counts = df["Manager"].value_counts()
    monthly = df.groupby("Month", as_index=False).size().rename(columns={"size": "Joinings"})
    monthly = monthly.sort_values("Month")

    top_recruiter = top_recruiter_counts.index[0]
    top_recruiter_n = int(top_recruiter_counts.iloc[0])
    top_recruiter_share = (top_recruiter_n / total) * 100

    top_client = top_client_counts.index[0]
    top_client_n = int(top_client_counts.iloc[0])
    top_client_share = (top_client_n / total) * 100

    us_share = (df["Work_Authorization"].eq("US CITIZEN").mean() * 100) if total else 0
    hybrid_share = (df["Work_Mode"].eq("Hybrid/Remote").mean() * 100) if total else 0

    if len(monthly) >= 2:
        prev_val = int(monthly.iloc[-2]["Joinings"])
        curr_val = int(monthly.iloc[-1]["Joinings"])
        if curr_val > prev_val:
            trend_line = f"Momentum is improving: latest month ({curr_val}) is above previous month ({prev_val})."
        elif curr_val < prev_val:
            trend_line = f"Momentum softened: latest month ({curr_val}) is below previous month ({prev_val})."
        else:
            trend_line = f"Momentum is stable: latest month and previous month are both {curr_val}."
    else:
        trend_line = "Only one month available after filtering, so trend comparison is limited."

    manager_variants = detect_name_variants(df["Manager"])
    if manager_variants:
        variant_note = "Potential manager-name variants found: " + "; ".join(
            [", ".join(group) for group in manager_variants[:3]]
        )
    else:
        variant_note = "No obvious manager-name spelling variants detected."

    insights = [
        f"Total joinings: {total} across {unique_clients} clients and {unique_recruiters} recruiters.",
        f"Top recruiter: {top_recruiter} ({top_recruiter_n} joinings, {top_recruiter_share:.1f}% share).",
        f"Top client: {top_client} ({top_client_n} joinings, {top_client_share:.1f}% share).",
        f"Top manager by load: {top_manager_counts.index[0]} ({int(top_manager_counts.iloc[0])} joinings).",
        trend_line,
        f"Work authorization mix: US Citizen share is {us_share:.1f}%.",
        f"Work model mix: Hybrid/Remote share is {hybrid_share:.1f}%.",
        variant_note,
        "Recommendation: standardize manager/recruiter naming and track client concentration monthly to de-risk dependency on a few sources.",
    ]
    return insights


def build_ppt(insights: List[str], kpis: Dict[str, str]) -> BytesIO:
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Recruitment Dashboard Summary"
    s1.placeholders[1].text = f"Generated on {datetime.now().strftime('%d %b %Y %H:%M')}"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "KPI Snapshot"
    tf = s2.shapes.placeholders[1].text_frame
    tf.clear()
    first = True
    for k, v in kpis.items():
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text = f"{k}: {v}"
        p.level = 0
        first = False

    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Executive Insights"
    tf2 = s3.shapes.placeholders[1].text_frame
    tf2.clear()
    first = True
    for line in insights:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        p.text = line
        p.level = 0
        first = False

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out


def main() -> None:
    st.set_page_config(page_title="Recruitment Executive Dashboard", layout="wide")

    # ── Theme selector (sidebar) ──────────────────────────────────────────
    theme_labels = ThemeFactory.all_labels()
    theme_key = st.sidebar.selectbox(
        "Theme",
        options=list(theme_labels.keys()),
        format_func=lambda k: theme_labels[k],
        index=1,  # default: corporate
    )
    factory = ThemeFactory(theme_key)
    theme = factory.theme

    # ── Inject CSS ────────────────────────────────────────────────────────
    st.markdown(factory.build_css(), unsafe_allow_html=True)

    # ── Algorithmic art header banner ─────────────────────────────────────
    render_particle_network(
        palette=theme.art_palette,
        height=120,
        n_particles=50,
        key=f"header_art_{theme_key}",
    )

    st.title("Recruitment Executive Dashboard")
    st.caption("Senior-analyst view · interactive filters · canvas art · themed exports")

    # ── Data loading ──────────────────────────────────────────────────────
    uploaded = st.file_uploader("Upload CSV (optional)", type=["csv"])
    source_name = "clean_recruitment_data.csv"

    try:
        raw_df = pd.read_csv(uploaded if uploaded is not None else source_name)
    except FileNotFoundError:
        st.error("File not found. Place `clean_recruitment_data.csv` in the same folder or upload it.")
        st.stop()
    except Exception as e:
        st.error(f"Could not read CSV: {e}")
        st.stop()

    try:
        df = preprocess(raw_df)
    except Exception as e:
        st.error(str(e))
        st.stop()

    if df.empty:
        st.warning("No valid rows left after preprocessing (check Joining_Date values).")
        st.stop()

    filtered = apply_filters(df)

    if filtered.empty:
        st.warning("No records found for selected filters.")
        st.stop()

    # ── KPIs ──────────────────────────────────────────────────────────────
    total_joinings = len(filtered)
    unique_candidates = filtered["Candidate"].nunique()
    unique_clients = filtered["Client"].nunique()
    unique_recruiters = filtered["Recruiter"].nunique()
    min_date = filtered["Joining_Date"].min().date()
    max_date = filtered["Joining_Date"].max().date()

    # Canvas KPI rings (top row)
    st.subheader("Key Performance Indicators")
    ring_cols = st.columns(4)
    ring_data = [
        (total_joinings, max(total_joinings, 1), "Total Joinings"),
        (unique_candidates, max(total_joinings, 1), "Unique Candidates"),
        (unique_clients, max(unique_clients, 1), "Unique Clients"),
        (unique_recruiters, max(unique_recruiters, 1), "Active Recruiters"),
    ]
    for col, (val, tot, lbl) in zip(ring_cols, ring_data):
        with col:
            render_kpi_ring(
                value=val,
                total=tot,
                label=lbl,
                accent=theme.css_vars["--accent-1"],
                bg_color=theme.css_vars["--bg-card"],
                size=150,
                key=f"ring_{lbl}_{theme_key}",
            )

    # Streamlit metrics below rings for accessibility
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Total Joinings", f"{total_joinings}")
    m2.metric("Unique Candidates", f"{unique_candidates}")
    m3.metric("Unique Clients", f"{unique_clients}")
    m4.metric("Active Recruiters", f"{unique_recruiters}")
    st.caption(f"Analysis Window: {min_date} to {max_date}")

    # ── Wave divider ──────────────────────────────────────────────────────
    render_wave_divider(
        color1=theme.css_vars["--accent-1"],
        color2=theme.css_vars["--accent-2"],
        key=f"wave_{theme_key}",
    )

    # ── Charts ────────────────────────────────────────────────────────────
    pc = factory.plotly_colors()

    monthly = (
        filtered.groupby("Month", as_index=False)
        .size()
        .rename(columns={"size": "Joinings"})
        .sort_values("Month")
    )
    top_clients = filtered["Client"].value_counts().head(10).reset_index()
    top_clients.columns = ["Client", "Joinings"]
    top_recruiters = filtered["Recruiter"].value_counts().head(10).reset_index()
    top_recruiters.columns = ["Recruiter", "Joinings"]
    manager_load = filtered["Manager"].value_counts().head(10).reset_index()
    manager_load.columns = ["Manager", "Joinings"]
    auth_mix = filtered["Work_Authorization"].value_counts().reset_index()
    auth_mix.columns = ["Work_Authorization", "Count"]
    mode_mix = filtered["Work_Mode"].value_counts().reset_index()
    mode_mix.columns = ["Work_Mode", "Count"]
    weekday = (
        filtered["Weekday"]
        .value_counts()
        .reindex(
            ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"],
            fill_value=0,
        )
        .reset_index()
    )
    weekday.columns = ["Weekday", "Joinings"]

    left, right = st.columns(2)

    with left:
        fig1 = px.line(
            monthly, x="Month", y="Joinings", markers=True,
            title="Monthly Joining Trend",
            color_discrete_sequence=pc["color_discrete_sequence"],
            template=pc["template"],
        )
        st.plotly_chart(fig1, use_container_width=True)

        fig3 = px.bar(
            top_recruiters.sort_values("Joinings"),
            x="Joinings", y="Recruiter", orientation="h",
            title="Recruiter Performance (Top 10)",
            color="Joinings",
            color_continuous_scale=pc["color_continuous_scale"],
            template=pc["template"],
        )
        fig3.update_layout(coloraxis_showscale=False)
        st.plotly_chart(fig3, use_container_width=True)

        fig5 = px.pie(
            auth_mix, names="Work_Authorization", values="Count",
            hole=0.55, title="Work Authorization Mix",
            color_discrete_sequence=pc["color_discrete_sequence"],
            template=pc["template"],
        )
        st.plotly_chart(fig5, use_container_width=True)

    with right:
        fig2 = px.bar(
            top_clients.sort_values("Joinings"),
            x="Joinings", y="Client", orientation="h",
            title="Top Clients by Joinings (Top 10)",
            color="Joinings",
            color_continuous_scale=pc["color_continuous_scale"],
            template=pc["template"],
        )
        fig2.update_layout(coloraxis_showscale=False)
        st.plotly_chart(fig2, use_container_width=True)

        fig4 = px.bar(
            manager_load.sort_values("Joinings"),
            x="Joinings", y="Manager", orientation="h",
            title="Manager Load Distribution (Top 10)",
            color="Joinings",
            color_continuous_scale=pc["color_continuous_scale"],
            template=pc["template"],
        )
        fig4.update_layout(coloraxis_showscale=False)
        st.plotly_chart(fig4, use_container_width=True)

        fig6 = px.bar(
            mode_mix, x="Work_Mode", y="Count",
            title="Work Mode Split",
            color="Work_Mode",
            color_discrete_sequence=pc["color_discrete_sequence"],
            template=pc["template"],
        )
        fig6.update_layout(showlegend=False)
        st.plotly_chart(fig6, use_container_width=True)

    fig7 = px.bar(
        weekday, x="Weekday", y="Joinings",
        title="Joinings by Weekday",
        color="Joinings",
        color_continuous_scale=pc["color_continuous_scale"],
        template=pc["template"],
    )
    fig7.update_layout(coloraxis_showscale=False)
    st.plotly_chart(fig7, use_container_width=True)

    # ── Mosaic art divider ────────────────────────────────────────────────
    render_mosaic_art(
        palette=theme.art_palette,
        height=72,
        key=f"mosaic_{theme_key}",
    )

    # ── Data quality ──────────────────────────────────────────────────────
    st.subheader("Data Quality Check")
    missing_counts = raw_df[REQUIRED_COLUMNS].isna().sum().reset_index()
    missing_counts.columns = ["Column", "Missing_Count"]
    st.dataframe(missing_counts, use_container_width=True)

    # ── Insights ──────────────────────────────────────────────────────────
    insights = generate_insights(filtered)
    st.subheader("Executive Insights (Slide-Ready)")
    for i, line in enumerate(insights, start=1):
        st.markdown(f"{i}. {line}")

    # ── Exports ───────────────────────────────────────────────────────────
    st.subheader("Exports")

    kpis = {
        "Total Joinings": str(total_joinings),
        "Unique Candidates": str(unique_candidates),
        "Unique Clients": str(unique_clients),
        "Active Recruiters": str(unique_recruiters),
        "Date Range": f"{min_date} to {max_date}",
    }

    insights_md = "# Recruitment Executive Insights\n\n" + "\n".join([f"- {i}" for i in insights])

    export_cols = st.columns(4)

    with export_cols[0]:
        st.download_button(
            "Download Insights (.md)",
            data=insights_md.encode("utf-8"),
            file_name=f"recruitment_insights_{datetime.now().strftime('%Y%m%d_%H%M')}.md",
            mime="text/markdown",
        )

    with export_cols[1]:
        st.download_button(
            "Download Filtered Data (.csv)",
            data=filtered.to_csv(index=False).encode("utf-8"),
            file_name=f"filtered_recruitment_data_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
            mime="text/csv",
        )

    with export_cols[2]:
        # Web Artifact — full HTML dashboard
        html_report = generate_html_report(
            filtered=filtered,
            insights=insights,
            kpis=kpis,
            theme_css_vars=theme.css_vars,
            accents=theme.accents,
            continuous_scale=theme.continuous_scale,
        )
        st.download_button(
            "Download Web Artifact (.html)",
            data=html_report.encode("utf-8"),
            file_name=f"recruitment_dashboard_{datetime.now().strftime('%Y%m%d_%H%M')}.html",
            mime="text/html",
        )

    with export_cols[3]:
        # Web Artifact — single chart (monthly trend)
        chart_html = generate_chart_artifact(
            fig=fig1,
            chart_title="Monthly Joining Trend",
            theme_css_vars=theme.css_vars,
        )
        st.download_button(
            "Download Chart Artifact (.html)",
            data=chart_html.encode("utf-8"),
            file_name=f"monthly_trend_{datetime.now().strftime('%Y%m%d_%H%M')}.html",
            mime="text/html",
        )

    if PPTX_AVAILABLE:
        ppt_buffer = build_ppt(insights, kpis)
        st.download_button(
            "Download Summary Deck (.pptx)",
            data=ppt_buffer,
            file_name=f"recruitment_summary_{datetime.now().strftime('%Y%m%d_%H%M')}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    else:
        st.info("Install `python-pptx` for PPT export.")


if __name__ == "__main__":
    main()
